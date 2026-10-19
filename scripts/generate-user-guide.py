import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "福驾科技CloudGuard方案使用说明书.md"
OUTPUT = ROOT / "docs" / "福驾科技CloudGuard方案使用说明书.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    paragraph.add_run(" 页")


def add_inline(paragraph, text):
    text = text.replace("  ", " ").strip()
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(33, 73, 104)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_screenshot(document, image_path, caption):
    full_path = SOURCE.parent / image_path
    image = DocxImage.from_file(str(full_path))
    width_px, height_px = image.px_width, image.px_height
    max_width, max_height = 6.5, 8.2
    ratio = width_px / height_px
    width = min(max_width, max_height * ratio)
    height = width / ratio
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(full_path), width=Inches(width), height=Inches(height))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.style = document.styles["Caption"]


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, "173F5F")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value.replace("`", "")
    document.add_paragraph()


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header.paragraphs[0].text = "福驾科技 CloudGuard｜安全治理方案使用说明书"
    section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor(16, 52, 77)
    styles["Title"].font.size = Pt(28)
    styles["Heading 1"].font.size = Pt(20)
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.size = Pt(12)


def build():
    document = Document()
    configure_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines = []
    index = 0
    image_number = 1

    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="No Spacing")
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(8)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line:
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line)
        if image_match:
            add_screenshot(document, image_match.group(2), f"图 {image_number}  {image_match.group(1)}")
            image_number += 1
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|(?:\s*:?-+:?\s*\|)+$", lines[index + 1]):
            rows = []
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1 and not document.paragraphs:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text)
                document.add_paragraph("AWS 原生安全治理与可重复交付平台").alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_heading(text, level=level)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        elif re.match(r"^\d+\.\s", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s", "", line))
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
        index += 1

    properties = document.core_properties
    properties.title = "福驾科技 CloudGuard 方案使用说明书"
    properties.subject = "AWS SMB 安全治理平台部署、运行与验收"
    properties.author = "福驾科技 / 博思云为"
    properties.keywords = "AWS, CloudGuard, SMB, Security, Governance, Config, Evidence"
    document.save(OUTPUT)
    print(f"Generated {OUTPUT}")


if __name__ == "__main__":
    build()
